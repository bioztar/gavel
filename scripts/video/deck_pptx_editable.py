"""The /architecture deck as an editable PPTX — real text boxes and shapes, not pictures.

Geometry and styling are read off the rendered page rather than retyped, so the layout
matches the deck and every string stays editable in PowerPoint. Brand glyphs are the one
thing that stays an image: they are inline SVG paths with no text to edit.
"""
from __future__ import annotations
import pathlib

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

OUT = pathlib.Path("/Users/alex/DEV/gavel/docs/gavel-architecture.pptx")
URL = "https://gavel.pro7ocol.com/architecture"
EXE = ("/Users/alex/Library/Caches/ms-playwright/chromium-1217/chrome-mac-arm64/"
       "Google Chrome for Testing.app/Contents/MacOS/Google Chrome for Testing")
VW, VH = 1920, 1080
SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)
PX = SLIDE_W / VW                       # one CSS pixel in EMU
NOTES = [
    "Karen sits in the call. Hear: Discord, SLNG speech-to-text, Postgres and Redis. "
    "Think: rules four times a second, Mastra orchestration, Nebius for the words. "
    "Speak: SLNG back into the call, fal moves her face, Vonage for the stage.",
    "Built and checked by machines too. Devin wrote PRs, Quality Clouds scanned the repo, "
    "Galtea scored the chair, Langfuse traced every model call.",
    "Discord was the demo. Teams is the product. Company rules set once, and a hand-off "
    "plugin that pushes the transcript, decisions and open items into the team's own agent.",
]

# Pulls every card and every run of text with its box, colour and weight.
SCRAPE = """
() => {
  const px = v => parseFloat(v) || 0;
  const rgb = v => {
    const m = (v || "").match(/\\d+/g);
    return m && m.length >= 3 ? [(+m[0]), (+m[1]), (+m[2]), m[3] === undefined ? 1 : +m[3]] : null;
  };
  const slide = document.querySelector(".slide.on");
  const box = el => { const r = el.getBoundingClientRect();
                      return {x: r.x, y: r.y, w: r.width, h: r.height}; };
  const cards = [...slide.querySelectorAll(".lane, .card, .tile, .todo, .stop, .plats > .chip")]
    .map(el => { const cs = getComputedStyle(el);
      return {...box(el), fill: rgb(cs.backgroundColor), line: rgb(cs.borderTopColor),
              lineW: px(cs.borderTopWidth), radius: px(cs.borderTopLeftRadius)}; });

  // A text block is an element whose own children are all inline - so a heading with an
  // <em> inside stays one box with two runs, instead of shattering into fragments.
  const inlineOnly = el => [...el.children].every(c => {
    const d = getComputedStyle(c).display;
    return d === "inline" || d === "inline-block";
  });
  const hasText = el => el.textContent.replace(/\\s+/g, " ").trim().length > 0;
  const blocks = [];
  const taken = [];
  for (const el of slide.querySelectorAll("*")) {
    if (taken.some(t => t.contains(el))) continue;
    const cs = getComputedStyle(el);
    if (cs.display === "none" || cs.visibility === "hidden") continue;
    if (el.tagName === "SVG" || el.closest("svg")) continue;
    if (!hasText(el) || !inlineOnly(el)) continue;
    if (el.classList.contains("mono")) continue;          // the square logo tiles
    const runs = [];
    for (const n of el.childNodes) {
      const t = (n.textContent || "").replace(/\\s+/g, " ");
      if (!t.trim()) continue;
      const style = n.nodeType === 1 ? getComputedStyle(n) : cs;
      runs.push({text: t, color: rgb(style.color),
                 bold: (parseInt(style.fontWeight) || 400) >= 600});
    }
    if (!runs.length) continue;
    taken.push(el);
    blocks.push({...box(el), runs, size: px(cs.fontSize),
                 lineHeight: px(cs.lineHeight) || px(cs.fontSize) * 1.3,
                 spacing: px(cs.letterSpacing), upper: cs.textTransform === "uppercase",
                 align: cs.textAlign, color: rgb(cs.color),
                 bold: (parseInt(cs.fontWeight) || 400) >= 600});
  }
  const glyphs = [...slide.querySelectorAll("svg, .mono, .face img")].map(box);
  return {cards, texts: blocks, glyphs};
}
"""


def solid(shape, rgba):
    if rgba and rgba[3] > 0.05:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*rgba[:3])
    else:
        shape.fill.background()


def render(prs, data, glyph_png, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xFB, 0xFB, 0xFD)

    for c in data["cards"]:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(c["x"] * PX)),
                                     Emu(int(c["y"] * PX)), Emu(int(c["w"] * PX)),
                                     Emu(int(c["h"] * PX)))
        solid(shp, c["fill"])
        if c["lineW"] and c["line"] and c["line"][3] > 0.05:
            shp.line.color.rgb = RGBColor(*c["line"][:3])
            shp.line.width = Pt(max(c["lineW"] * 0.75, 0.75))
        else:
            shp.line.fill.background()
        shp.adjustments[0] = min(0.5, (c["radius"] / min(c["w"], c["h"])) if c["w"] and c["h"] else 0.1)
        shp.shadow.inherit = False

    # the brand marks are inline SVG paths; keep them as one transparent overlay image
    slide.shapes.add_picture(str(glyph_png), 0, 0, width=prs.slide_width, height=prs.slide_height)

    for t in data["texts"]:
        # PowerPoint's metrics are not the browser's, so a box measured to the pixel
        # reflows. One-line labels get room and are told never to wrap; paragraphs keep
        # wrapping but are clamped to the slide so they cannot run off the right edge.
        one_line = t["h"] <= t["lineHeight"] * 1.4
        want = t["w"] * (1.9 if one_line else 1.04) + t["size"] * 2
        width = min(want, VW - t["x"] - 40)
        tb = slide.shapes.add_textbox(Emu(int(t["x"] * PX)), Emu(int(t["y"] * PX)),
                                      Emu(int(width * PX)),
                                      Emu(int((t["h"] + t["size"] * 0.7) * PX)))
        tf = tb.text_frame
        tf.word_wrap = not one_line
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(t["align"],
                                                                               PP_ALIGN.LEFT)
        p.line_spacing = max(t["lineHeight"] / t["size"], 1.0) if t["size"] else 1.15
        for r in t["runs"]:
            run = p.add_run()
            run.text = r["text"].upper() if t["upper"] else r["text"]
            f = run.font
            f.size = Pt(round(t["size"] * 0.70, 1))
            f.bold = r["bold"]
            f.name = "Helvetica Neue"
            if t["spacing"]:
                run.font._rPr.set("spc", str(int(t["spacing"] * 100)))
            if r["color"]:
                f.color.rgb = RGBColor(*r["color"][:3])
    slide.notes_slide.notes_text_frame.text = notes


if __name__ == "__main__":
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    tmp = pathlib.Path("/Users/alex/DEV/_assets/gavel-video/shots")

    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=EXE)
        page = b.new_page(viewport={"width": VW, "height": VH}, device_scale_factor=2)
        page.goto(URL, wait_until="networkidle")
        page.wait_for_timeout(1200)
        for i in range(3):
            if i:
                page.keyboard.press("ArrowRight")
                page.wait_for_timeout(900)
            page.evaluate("""() => {
                const d = document.getElementById('dots'); if (d) d.style.display = 'none';
                const h = document.querySelector('.hint'); if (h) h.style.display = 'none';
            }""")
            data = page.evaluate(SCRAPE)
            # Glyph-only pass. Hiding the text would collapse the flex layout and move the
            # logos, so the text is made invisible instead and everything stays put.
            page.evaluate("""() => {
                const st = document.createElement('style');
                st.id = 'glyphpass';
                st.textContent = `.slide.on, .slide.on * { color: transparent !important;
                    -webkit-text-fill-color: transparent !important; }
                  .slide.on .lane, .slide.on .card, .slide.on .tile, .slide.on .todo,
                  .slide.on .stop { background: transparent !important; border-color:
                    transparent !important; box-shadow: none !important; }
                  .slide.on .lane::before, .slide.on .return i { opacity: 0 !important; }
                  html, body { background: transparent !important; background-image: none !important; }`;
                document.head.append(st);
            }""")
            page.wait_for_timeout(250)
            glyph = tmp / f"deck-glyphs-{i + 1}.png"
            page.screenshot(path=str(glyph), omit_background=True)
            render(prs, data, glyph, NOTES[i])
            print(f"slide {i + 1}: {len(data['cards'])} shapes, {len(data['texts'])} text boxes")
            page.reload(wait_until="networkidle")
            page.wait_for_timeout(700)
            for _ in range(i):
                page.keyboard.press("ArrowRight")
                page.wait_for_timeout(500)
        b.close()

    prs.save(OUT)
    print(f"\n{OUT}  {OUT.stat().st_size // 1024} KB")
