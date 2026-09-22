"""The live /architecture deck as a PPTX — one full-bleed slide per HTML slide."""
from __future__ import annotations
import pathlib
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

OUT = pathlib.Path("/Users/alex/DEV/gavel/docs/gavel-architecture.pptx")
SHOTS = pathlib.Path("/Users/alex/DEV/_assets/gavel-video/shots")
URL = "https://gavel.pro7ocol.com/architecture"
EXE = ("/Users/alex/Library/Caches/ms-playwright/chromium-1217/chrome-mac-arm64/"
       "Google Chrome for Testing.app/Contents/MacOS/Google Chrome for Testing")
NOTES = [
    "Karen sits in the call. Hear: Discord, SLNG speech-to-text, Postgres and Redis. "
    "Think: rules four times a second, Mastra orchestration, Nebius for the words. "
    "Speak: SLNG back into the call.",
    "Built and checked by machines too. Devin wrote PRs, Quality Clouds scanned the repo, "
    "Galtea scored the chair across five dimensions, Langfuse traced every model call.",
    "Discord was the demo. Teams is the product. One container knows what a voice call is. "
    "Company rules set once and enforced everywhere, and a hand-off plugin that pushes the "
    "transcript, decisions and open items into the team's own knowledge agent.",
]

pngs = []
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=EXE)
    page = b.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=2)
    page.goto(URL, wait_until="networkidle")
    page.wait_for_timeout(1200)
    for i in range(3):
        if i:
            page.keyboard.press("ArrowRight")
            page.wait_for_timeout(900)
        # the nav dots and the arrow hint belong to the web deck, not to a slide
        page.evaluate("""() => {
            const d = document.getElementById('dots'); if (d) d.style.display = 'none';
            const h = document.querySelector('.hint'); if (h) h.style.display = 'none';
        }""")
        page.wait_for_timeout(200)
        f = SHOTS / f"deck-{i + 1}.png"
        page.screenshot(path=str(f))
        pngs.append(f)
        print("shot", f.name)
    b.close()

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9, 13.333 x 7.5 in
for png, note in zip(pngs, NOTES):
    slide = prs.slides.add_slide(prs.slide_layouts[6])            # blank
    slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text = note
prs.save(OUT)
print(f"\n{OUT}  {OUT.stat().st_size // 1024} KB, {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
